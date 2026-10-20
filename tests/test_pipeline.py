import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from nsw_crime_trends.flags import build_flags, level_shifts
from nsw_crime_trends.io import load_bocsar, load_erp_june
from nsw_crime_trends.rates import add_rates, annual_counts, fixed_denominator_bias
from nsw_crime_trends.trends import kendall_60m, rate_trend_qp

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data_raw"


def make_xlsx(path, months, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["State", "Offence category", "Subcategory", "2025 population", "2026 population", *months])
    for r in rows:
        ws.append(r)
    wb.save(path)


def make_erp(path, years, value=1_000_000):
    lines = ["DATAFLOW,MEASURE,SEX,AGE,REGION,FREQ,TIME_PERIOD,OBS_VALUE,UNIT_MEASURE,OBS_STATUS,OBS_COMMENT"]
    for y in years:
        for q in (1, 2, 3, 4):
            v = value + (5 if q != 2 else 0)
            lines.append(f"ABS:ERP_Q(1.0.0),1,3,TOT,1,Q,{y}-Q{q},{v},PSNS,,")
        lines.append(f"ABS:ERP_Q(1.0.0),1,3,TOT,2,Q,{y}-Q2,999,PSNS,,")  # another state, must be ignored
    Path(path).write_text("\n".join(lines))


def test_load_bocsar_units_and_shape(tmp_path):
    months = [dt.datetime(2024, m, 1) for m in range(1, 13)]
    make_xlsx(tmp_path / "x.xlsx", months, [
        ["NSW", "Homicide", "Murder *", 1, 1, *[1] * 12],
        ["NSW", "Arson", None, 1, 1, *[2] * 12],
    ])
    df = load_bocsar(tmp_path / "x.xlsx")
    assert set(df["offence"]) == {"Homicide: Murder", "Arson"}
    assert df.loc[df["offence"] == "Homicide: Murder", "unit"].unique().tolist() == ["victims"]
    assert df.loc[df["offence"] == "Arson", "unit"].unique().tolist() == ["incidents"]
    assert len(df) == 24 and df["count"].sum() == 36


def test_load_bocsar_rejects_unexpected_header(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"; ws.append(["a", "b", "c"]); wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError):
        load_bocsar(tmp_path / "bad.xlsx")


def test_erp_uses_june_quarter_and_nsw_only(tmp_path):
    make_erp(tmp_path / "e.csv", [2020, 2021], value=500)
    s = load_erp_june(tmp_path / "e.csv")
    assert s.to_dict() == {2020: 500, 2021: 500}


def test_annual_counts_drop_partial_year_and_rates():
    months = pd.date_range("2024-01-01", "2025-03-01", freq="MS")
    m = pd.DataFrame({"offence": "X", "category": "X", "unit": "incidents", "month": months, "count": 10})
    a = annual_counts(m)
    assert a["year"].tolist() == [2024] and a["count"].tolist() == [120]
    r = add_rates(a, pd.Series({2024: 200_000}))
    assert r["rate_per_100k"].iloc[0] == pytest.approx(60.0)
    with pytest.raises(ValueError):
        add_rates(a, pd.Series({2023: 1}))


def test_fixed_denominator_understates_when_population_grows():
    a = pd.DataFrame({"offence": ["X"], "year": [1995], "count": [100], "erp": [50_000], "rate_per_100k": [200.0]})
    b = fixed_denominator_bias(a, 100_000, 1995)
    assert b["rate_fixed_pop"].iloc[0] == pytest.approx(100.0)
    assert b["pct_understated"].iloc[0] == pytest.approx(-50.0)


def test_kendall_up_and_aapc():
    idx = pd.date_range("2020-01-01", periods=60, freq="MS")
    s = pd.Series(np.arange(60) + 100, index=idx)
    k = kendall_60m(s)
    assert k["tested"] and k["direction"] == "up" and k["tau_b"] == pytest.approx(1.0)
    t = s.values.reshape(5, 12).sum(axis=1)
    assert k["avg_annual_pct_change"] == pytest.approx(((t[-1] / t[0]) ** 0.25 - 1) * 100)


def test_kendall_not_tested_when_low_counts():
    idx = pd.date_range("2020-01-01", periods=60, freq="MS")
    k = kendall_60m(pd.Series(1, index=idx))
    assert k["tested"] is False and "below 20" in k["reason"]


def test_qp_recovers_known_rate_change():
    years = np.arange(2016, 2026)
    erp = 1_000_000 * 1.02 ** (years - 2016)               # population grows 2%/yr
    count = np.round(erp * 1e-3 * 1.05 ** (years - 2016))  # rate grows 5%/yr
    d = pd.DataFrame({"year": years, "count": count, "erp": erp})
    q = rate_trend_qp(d, 2016, 2025)
    assert q["fitted"] and q["apc"] == pytest.approx(5.0, abs=0.05)
    assert q["apc_lo"] < q["apc"] < q["apc_hi"]
    q2 = rate_trend_qp(d, 2016, 2025, exclude=(2020, 2021))
    assert q2["n_years"] == 8


def test_level_shift_and_new_series_flags():
    a = pd.Series([100, 102, 99, 101, 100, 250, 252, 249], index=range(2000, 2008))
    shifts = level_shifts(a)
    assert [s["year"] for s in shifts] == [2005]
    months = pd.date_range("2020-01-01", "2024-12-01", freq="MS")
    counts = [0] * 36 + [5] * 24
    m = pd.DataFrame({"offence": "New", "category": "New", "unit": "incidents", "month": months, "count": counts})
    fl = build_flags(m, annual_counts(m))
    assert any(f.startswith("NEW_SERIES: zero until 2023-01") for f in fl["flags"].iloc[0])


@pytest.mark.skipif(not (RAW / "Incident_by_NSW.xlsx").exists(), reason="public data not downloaded")
def test_real_data_end_to_end(tmp_path):
    from main import run
    s = run(RAW / "Incident_by_NSW.xlsx", RAW / "ABS_ERP_Q_NSW_persons.csv", tmp_path)
    assert s["coverage"]["n_offence_series"] == 62
    assert s["coverage"]["first_month"] == "1995-01"
    # independent re-count straight from the workbook for one offence and year
    wb = openpyxl.load_workbook(RAW / "Incident_by_NSW.xlsx", read_only=True, data_only=True)
    rows = list(wb["Data"].iter_rows(values_only=True))
    hdr = rows[0]
    row = next(r for r in rows[1:] if r[2] == "Domestic violence related assault")
    direct = sum(v for h, v in zip(hdr, row) if isinstance(h, dt.datetime) and h.year == 2025)
    ann = {r["year"]: r["count"] for r in s["annual"]["Assault: Domestic violence related assault"]}
    assert ann[2025] == direct
    murder = next(o for o in s["offences"] if o["offence"] == "Homicide: Murder")
    assert murder["unit"] == "victims"
    # Independently reconcile the denominator illustration to a single source
    # row, so unrelated incident or victim counts cannot enter its numerator.
    retail_row = next(r for r in rows[1:] if r[2] == "Steal from retail store")
    retail_1995 = sum(v for h, v in zip(hdr, retail_row)
                      if isinstance(h, dt.datetime) and h.year == 1995)
    dn = s["denominator_check"]
    assert dn["example_offence"] == "Theft: Steal from retail store"
    assert dn["example_unit"] == "incidents"
    assert dn["example_1995_count"] == retail_1995
    assert not any(k.startswith("all_offences") for k in dn)
    assert dn["example_1995_rate_with_1995_erp"] == round(retail_1995 / s["erp_june"][1995] * 1e5, 1)
    from build_sample import build
    page = build(s, "2026-10-20")
    assert f"{retail_1995:,} incidents were recorded in 1995" in page
    assert "All 62 offences together" not in page
    assert "PhD (biostatistics)" not in page
    assert (tmp_path / "trend_summary.csv").exists()


def test_series_year_with_missing_month_is_excluded_not_undercounted(tmp_path):
    months = [dt.datetime(2024, m, 1) for m in range(1, 13)] + [dt.datetime(2025, m, 1) for m in range(1, 13)]
    a_vals = [10] * 24
    b_vals = [10] * 24
    b_vals[15] = None  # Apr 2025 blank for offence B only
    make_xlsx(tmp_path / "m.xlsx", months, [["NSW", "A", None, 1, 1, *a_vals], ["NSW", "B", None, 1, 1, *b_vals]])
    monthly = load_bocsar(tmp_path / "m.xlsx")
    assert monthly["count"].isna().sum() == 1  # blank stays missing, not 0
    ann = annual_counts(monthly)
    got = {(r.offence, r.year): r.count for r in ann.itertuples()}
    assert got == {("A", 2024): 120, ("A", 2025): 120, ("B", 2024): 120}  # B-2025 excluded, not 110
    ex = ann.attrs["excluded_series_years"]
    assert ex == [{"offence": "B", "year": 2025, "n_months": 12, "n_present": 11}]
    fl = build_flags(monthly, ann)
    b_flags = fl.loc[fl["offence"] == "B", "flags"].iloc[0]
    assert any(f.startswith("MISSING_MONTHS: 1") for f in b_flags)
    idx = pd.date_range("2021-01-01", periods=60, freq="MS")
    s = pd.Series([50] * 60, index=idx, dtype="Int64"); s.iloc[10] = pd.NA
    assert kendall_60m(s)["tested"] is False
